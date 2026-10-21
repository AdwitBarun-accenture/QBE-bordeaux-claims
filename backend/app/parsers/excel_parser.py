import io

from openpyxl import load_workbook

from app.parsers.base import ParsedFile, UnsupportedFile
from app.parsers.grid import grid_to_records

PREMIUM_HINTS = ("prem", "gwp", "umr", "inception")
CLAIM_HINTS = ("claim", "loss", "reserve", "incurred")


def list_sheets(data: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]


def _pick_sheet(wb) -> str:
    """Choose the visible sheet that looks most like a bordereaux (most rows +
    bordereaux keywords in its first rows)."""
    best, best_score = None, -1
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        head = " ".join(str(c or "").lower() for row in ws.iter_rows(max_row=10, values_only=True) for c in row)
        kw = sum(head.count(k) for k in PREMIUM_HINTS + CLAIM_HINTS)
        score = kw * 100 + (ws.max_row or 0)
        if score > best_score:
            best, best_score = ws.title, score
    return best


def parse_excel(data: bytes, sheet: str | None = None) -> ParsedFile:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise UnsupportedFile(f"cannot open workbook: {e}") from e
    sheets = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    chosen = sheet if sheet in sheets else _pick_sheet(wb)
    ws = wb[chosen]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    headers, rows, h_idx, dropped = grid_to_records(grid)
    pf = ParsedFile(headers=headers, rows=rows, fmt="xlsx", sheet_name=chosen, sheets=sheets,
                    header_row_index=h_idx, dropped_empty_rows=dropped)
    if len(sheets) > 1:
        pf.warnings.append(f"workbook has {len(sheets)} sheets; auto-selected '{chosen}'")
    return pf
